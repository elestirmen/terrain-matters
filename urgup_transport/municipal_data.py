"""Readers for the files the municipality delivered, in the shapes they arrived in.

These are not canonical datasets and nothing here writes one. Each function turns
one delivered file into plain records, and refuses rather than guesses when the
file does not say something: a spreadsheet that records how many runs a line
makes does not record which days it runs, and inventing a day set here would put
a number in front of a passenger that nobody at the municipality ever stated.

The parsers live together because they share one property — the source is a
human-maintained office document, so the reading has to survive a merged header,
a name Excel decided was a date, and a column total that does not add up. Each of
those is handled explicitly and reported, never silently corrected.
"""

from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

KML_NS = {"k": "http://www.opengis.net/kml/2.2"}


class MunicipalDataError(ValueError):
    """Raised when a delivered file cannot be read as the format it claims."""


# ── KML ──────────────────────────────────────────────────────────────────
@dataclass(frozen=True)
class KmlFeature:
    """One placemark, with the folder chain that gave it its meaning.

    The folder is not decoration: in `GÜZERGAHLAR.kml` the line's identity is the
    folder name and most placemarks inside it are unnamed fragments, so a reader
    that kept only `<name>` would lose which line a geometry belongs to.
    """

    name: str
    folder_path: tuple[str, ...]
    geometry: dict[str, Any]

    @property
    def folder(self) -> str:
        return self.folder_path[-1] if self.folder_path else ""


def _text(node: ET.Element | None, path: str) -> str:
    if node is None:
        return ""
    child = node.find(path, KML_NS)
    if child is None or child.text is None:
        return ""
    return child.text.strip()


def _tag(element: ET.Element) -> str:
    return element.tag.split("}", 1)[-1]


def _coordinates(raw: str) -> list[tuple[float, float]]:
    """Longitude/latitude pairs, dropping the altitude Netcad writes as a third value.

    Altitude is dropped rather than carried because the draft contract is
    explicitly two-dimensional (`flatten_to_2d`), and height is derived from the
    DEM at read time. A z that survived to here would be a second, staler answer.
    """
    points: list[tuple[float, float]] = []
    for token in raw.split():
        parts = token.split(",")
        if len(parts) < 2:
            continue
        try:
            points.append((float(parts[0]), float(parts[1])))
        except ValueError:
            continue
    return points


def _geometry(placemark: ET.Element) -> dict[str, Any] | None:
    point = _text(placemark, ".//k:Point/k:coordinates")
    if point:
        coordinates = _coordinates(point)
        if coordinates:
            return {"type": "Point", "coordinates": list(coordinates[0])}

    line = _text(placemark, ".//k:LineString/k:coordinates")
    if line:
        coordinates = _coordinates(line)
        if len(coordinates) >= 2:
            return {"type": "LineString", "coordinates": [list(point) for point in coordinates]}

    polygon = placemark.find(".//k:Polygon", KML_NS)
    if polygon is not None:
        outer = _coordinates(_text(polygon, ".//k:outerBoundaryIs//k:coordinates"))
        if len(outer) >= 3:
            rings = [_closed_ring(outer)]
            for inner in polygon.findall(".//k:innerBoundaryIs//k:LinearRing", KML_NS):
                hole = _coordinates(_text(inner, "k:coordinates"))
                if len(hole) >= 3:
                    rings.append(_closed_ring(hole))
            return {"type": "Polygon", "coordinates": rings}
    return None


def _closed_ring(points: list[tuple[float, float]]) -> list[list[float]]:
    ring = [list(point) for point in points]
    if ring[0] != ring[-1]:
        ring.append(list(ring[0]))
    return ring


def read_kml(path: Path) -> list[KmlFeature]:
    """Every placemark that carries a geometry, in document order."""
    try:
        root = ET.parse(path).getroot()
    except (OSError, ET.ParseError) as exc:
        raise MunicipalDataError(f"KML okunamadı ({path.name}): {exc}") from exc

    features: list[KmlFeature] = []

    def walk(container: ET.Element, chain: tuple[str, ...]) -> None:
        for child in container:
            tag = _tag(child)
            if tag in {"Document", "Folder"}:
                name = _text(child, "k:name")
                walk(child, (*chain, name) if name else chain)
            elif tag == "Placemark":
                geometry = _geometry(child)
                if geometry is not None:
                    features.append(
                        KmlFeature(
                            name=_text(child, "k:name"),
                            folder_path=chain,
                            geometry=geometry,
                        )
                    )

    walk(root, ())
    return features


# ── declared service ─────────────────────────────────────────────────────
@dataclass(frozen=True)
class ServiceBand:
    """A named part of the day, and how many runs the municipality put in it.

    The source counts runs; it does not state a headway. `headway_minutes`
    converts one into the other on the only assumption the sheet supports — that
    the runs are spread evenly across the band.
    """

    label: str
    start_seconds: int
    end_seconds: int
    trips: int

    @property
    def duration_minutes(self) -> float:
        return (self.end_seconds - self.start_seconds) / 60.0

    @property
    def headway_minutes(self) -> float:
        """The even spacing implied by the run count, rounded up to six seconds.

        Rounded up, not to nearest: GTFS reads a window as departures at
        `start`, `start + headway`, … while below `end`, so a headway rounded
        down fits one more departure into the band than the municipality
        declared. Rounding up can only ever promise fewer runs than were stated,
        which is the safe direction to be wrong in.
        """
        exact = self.duration_minutes / self.trips
        return math.ceil(exact * 10.0) / 10.0


@dataclass(frozen=True)
class ServiceLine:
    """One line's declared operation, and what the sheet left unsaid."""

    name: str
    bands: tuple[ServiceBand, ...]
    vehicles: int | None = None
    declared_daily_trips: int | None = None
    round_trip_text: str = ""
    round_trip_minutes: float | None = None
    notes: tuple[str, ...] = ()

    @property
    def band_trips(self) -> int:
        return sum(band.trips for band in self.bands)

    @property
    def publishable(self) -> bool:
        return bool(self.bands)


# The bands are named by the header cells of `DURAK SERVİS SAATLERİ.xlsx`, which
# carry their own clock range. They are read from the file rather than hardcoded;
# these are only the labels a passenger sees.
_BAND_HEADER = re.compile(r"^\s*SEFER\s+(?P<label>.+?)\s*\(\s*(?P<start>[\d:.]+)\s*-\s*(?P<end>[\d:.]+)")


def _clock_seconds(text: str) -> int | None:
    parts = re.split(r"[:.]", text.strip())
    if not parts or not all(part.isdigit() for part in parts):
        return None
    hours = int(parts[0])
    minutes = int(parts[1]) if len(parts) > 1 else 0
    if minutes > 59:
        return None
    return hours * 3600 + minutes * 60


def _cell_text(value: Any) -> str:
    """A header or name cell as text, undoing Excel's date guess.

    `15 TEMMUZ` is a neighbourhood here and a date to Excel, which stored it as
    2025-07-15. Left alone it would reach the feed as a line called "2025-07-15
    00:00:00"; read back through the Turkish month name it is the line the rest
    of the sheet calls TOKİ.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        return f"{value.day} {_TURKISH_MONTHS[value.month]}"
    if isinstance(value, date):
        return f"{value.day} {_TURKISH_MONTHS[value.month]}"
    return str(value).strip()


_TURKISH_MONTHS = {
    1: "OCAK",
    2: "ŞUBAT",
    3: "MART",
    4: "NİSAN",
    5: "MAYIS",
    6: "HAZİRAN",
    7: "TEMMUZ",
    8: "AĞUSTOS",
    9: "EYLÜL",
    10: "EKİM",
    11: "KASIM",
    12: "ARALIK",
}


def _int_cell(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else None
    match = re.search(r"\d+", str(value))
    return int(match.group(0)) if match else None


def _minutes_cell(value: Any) -> float | None:
    """`22,2 DK`, `34-38 DK` or a real time cell, as minutes.

    A range is read as its lower bound: the sheet's own totals are built from the
    shorter figure, and a round trip quoted as a range is the operator saying the
    upper end is traffic, not schedule.
    """
    if isinstance(value, time):
        return value.hour * 60 + value.minute + value.second / 60.0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value or "").strip()
    if not text:
        return None
    match = re.search(r"\d+(?:[.,]\d+)?", text)
    if not match:
        return None
    return float(match.group(0).replace(",", "."))


def read_service_sheet(path: Path) -> list[ServiceLine]:
    """The per-line run counts from `DURAK SERVİS SAATLERİ.xlsx`.

    Two tables share the sheet: run counts per band at the top, driver shifts
    below. Only the top one is read here — the shift table repeats the totals and
    adds who drives them, which is rostering, not what a passenger is told.

    A line whose bands are all empty comes back with no bands rather than being
    dropped, so the caller can say *why* it cannot be published instead of the
    line quietly not existing.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exception types
        raise MunicipalDataError(f"Sefer saatleri dosyası okunamadı ({path.name}): {exc}") from exc

    try:
        sheet = workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    if not rows:
        raise MunicipalDataError("Sefer saatleri dosyası boş.")

    header = [_cell_text(value) for value in rows[0]]
    try:
        name_column = header.index("GÜZERGAH")
    except ValueError as exc:
        raise MunicipalDataError("Sefer saatleri dosyasında 'GÜZERGAH' sütunu yok.") from exc

    bands: list[tuple[int, str, int, int]] = []
    for index, cell in enumerate(header):
        match = _BAND_HEADER.match(cell)
        if not match:
            continue
        start = _clock_seconds(match.group("start"))
        end = _clock_seconds(match.group("end"))
        if start is None or end is None or end <= start:
            raise MunicipalDataError(f"Sefer bandı başlığı okunamadı: {cell!r}")
        bands.append((index, match.group("label").strip().title(), start, end))
    if not bands:
        raise MunicipalDataError("Sefer saatleri dosyasında saat bandı başlığı bulunamadı.")

    vehicle_column = header.index("TAHSİL EDİLEN ARAÇ SAYISI") if "TAHSİL EDİLEN ARAÇ SAYISI" in header else None
    total_column = header.index("TOPLAM GÜNLÜK SEFER") if "TOPLAM GÜNLÜK SEFER" in header else None
    round_trip_column = header.index("R. SÜRE") if "R. SÜRE" in header else None

    lines: list[ServiceLine] = []
    for row in rows[1:]:
        if not any(_cell_text(value) for value in row):
            # A blank row ends the table. The same sheet carries a second one
            # below it — which driver works which shift — whose first column is
            # also a line name. Read past the gap and every roster row arrives as
            # a phantom line, one of them called `00:22:02`.
            break

        name = _cell_text(row[name_column] if name_column < len(row) else None)
        if not name:
            # The sheet's own column-total row has no line name. It is skipped
            # rather than read: its afternoon figure is four runs short of the
            # column it sums, and the per-line rows all agree with their own
            # totals, so the row is arithmetic nobody depends on.
            continue

        row_bands = []
        for index, label, start, end in bands:
            trips = _int_cell(row[index]) if index < len(row) else None
            if trips:
                row_bands.append(ServiceBand(label=label, start_seconds=start, end_seconds=end, trips=trips))

        declared_total = _int_cell(row[total_column]) if total_column is not None and total_column < len(row) else None
        round_trip_raw = (
            row[round_trip_column]
            if round_trip_column is not None and round_trip_column < len(row)
            else None
        )
        vehicles = _int_cell(row[vehicle_column]) if vehicle_column is not None and vehicle_column < len(row) else None

        notes: list[str] = []
        if not row_bands:
            notes.append("Saat bandı başına sefer sayısı verilmemiş.")
        elif declared_total is not None and declared_total != sum(band.trips for band in row_bands):
            notes.append(
                f"Bantların toplamı {sum(band.trips for band in row_bands)}, "
                f"dosyadaki günlük toplam {declared_total}."
            )
        for band in row_bands:
            if band.trips == 1:
                notes.append(f"{band.label}: tek sefer var, saati belirtilmemiş.")

        lines.append(
            ServiceLine(
                name=name,
                bands=tuple(row_bands),
                vehicles=vehicles,
                declared_daily_trips=declared_total,
                round_trip_text=str(round_trip_raw or "").strip() if not isinstance(round_trip_raw, time) else "",
                round_trip_minutes=_minutes_cell(round_trip_raw),
                notes=tuple(notes),
            )
        )
    return lines


# ── population ───────────────────────────────────────────────────────────
@dataclass(frozen=True)
class PopulationRecord:
    """One MAKS neighbourhood statistic, age band by age band."""

    mahalle: str
    source_file: str
    bands: tuple[tuple[str, int, int], ...] = field(default=())
    male: int = 0
    female: int = 0
    total: int = 0


_MAKS_TITLE = re.compile(r"NEVŞEHİR\s+ÜRGÜP\s+(?P<name>.+?)\s+MAHALLES", re.IGNORECASE)


def read_maks_population(path: Path) -> PopulationRecord:
    """One `MAKS-*.xlsx` age table.

    Two of the delivered files are not spreadsheets at all: Excel saved them as
    an HTML frameset whose `_dosyalar/sheet001.htm` was never sent, so the file
    holds a page of script and no numbers. That is reported as a missing
    delivery rather than a parse failure, because the fix is to ask for the file
    again, not to read it harder.
    """
    head = path.read_bytes()[:512].lstrip()
    if head[:1] == b"<":
        raise MunicipalDataError(
            f"{path.name}: Excel değil, veri sayfası eksik gönderilmiş bir HTML çerçevesi "
            "(_dosyalar klasörü teslim edilmemiş). Dosyanın yeniden istenmesi gerekiyor."
        )

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise MunicipalDataError(f"{path.name}: nüfus dosyası okunamadı ({exc}).") from exc

    try:
        rows = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()

    name = ""
    for row in rows[:6]:
        for value in row:
            match = _MAKS_TITLE.search(str(value or ""))
            if match:
                name = match.group("name").strip()
                break
        if name:
            break
    if not name:
        raise MunicipalDataError(f"{path.name}: mahalle adı başlıkta bulunamadı.")

    bands: list[tuple[str, int, int]] = []
    male = female = total = 0
    for row in rows:
        label = str(row[0] or "").strip() if row else ""
        if not label or len(row) < 4:
            continue
        values = [_int_cell(value) for value in row[1:4]]
        if any(value is None for value in values):
            continue
        row_male, row_female, row_total = (int(value) for value in values)  # type: ignore[arg-type]
        if label.upper().startswith("TOPLAM"):
            male, female, total = row_male, row_female, row_total
        elif "YAŞ" in label.upper():
            bands.append((label, row_male, row_female))

    if not total:
        raise MunicipalDataError(f"{path.name}: 'Toplam' satırı bulunamadı.")
    band_total = sum(row_male + row_female for _, row_male, row_female in bands)
    if band_total != total:
        raise MunicipalDataError(
            f"{path.name}: yaş bantlarının toplamı {band_total}, dosyadaki toplam {total}."
        )

    return PopulationRecord(
        mahalle=name,
        source_file=path.name,
        bands=tuple(bands),
        male=male,
        female=female,
        total=total,
    )


# ── names ────────────────────────────────────────────────────────────────
# Shapefile attribute names are capped at ten characters, so the boundaries
# arrived truncated and, for the two Aksalur neighbourhoods, abbreviated as well.
# The published name is matched to an official one by normalizing both; these are
# the pairs normalization cannot reach on its own.
NAME_ALIASES = {
    "AK_AĞCAŞAR": "AKSALUR AĞCAŞAR",
    "AK_SALUR": "AKSALUR SALUR",
}


def normalize_name(value: str) -> str:
    """Fold a name to letters and digits, so `BAHÇELIEVL` can meet `BAHÇELİEVLER`.

    Turkish dotted and dotless i are folded together deliberately: the truncated
    names lost the dot when they passed through a shapefile, and refusing to
    match `FATIH` with `FATİH` would mean refusing every name in the file.
    """
    lowered = value.replace("İ", "i").replace("I", "ı").casefold()
    stripped = unicodedata.normalize("NFD", lowered)
    folded = "".join(char for char in stripped if not unicodedata.combining(char))
    for turkish, latin in (("ı", "i"), ("ş", "s"), ("ğ", "g"), ("ö", "o"), ("ü", "u"), ("ç", "c")):
        folded = folded.replace(turkish, latin)
    return re.sub(r"[^a-z0-9]", "", folded)


# Registry exports are shouted; a map is not. These stay upper because they are
# abbreviations rather than words — title-casing them produces a name nobody in
# Ürgüp writes.
UPPERCASE_NAMES = frozenset({"EVKA", "TOKİ", "TOKI", "KYK"})


def turkish_title(name: str) -> str:
    """`BAHÇELİEVLER` as `Bahçelievler`, with Turkish's two letter i.

    `str.title()` is wrong here in both directions: it turns `FATİH` into `Fati̇h`
    by decomposing the dotted capital, and `KAVAKLIONU` into `Kavaklionu` by
    lowering the dotless I to a dotted i. The mapping is done explicitly instead.
    """
    words = []
    for word in name.split():
        if word.upper() in UPPERCASE_NAMES:
            words.append(word.upper())
            continue
        lowered = word.replace("I", "ı").replace("İ", "i").lower()
        head = lowered[:1].replace("i", "İ").replace("ı", "I").upper()
        words.append(head + lowered[1:])
    return " ".join(words)


def official_name_for(published: str, candidates: dict[str, str]) -> str | None:
    """The official spelling of a published name, or None when nothing matches.

    Matching is by prefix in the truncation's direction only — a published name
    may be a shortened official one, never the other way round — so a genuinely
    new neighbourhood cannot be absorbed into an existing one by accident.
    """
    alias = NAME_ALIASES.get(published.strip().upper())
    if alias:
        for official in candidates.values():
            if normalize_name(official) == normalize_name(alias):
                return official
        return None

    key = normalize_name(published)
    if not key:
        return None
    exact = candidates.get(key)
    if exact:
        return exact
    matches = [official for candidate, official in candidates.items() if candidate.startswith(key)]
    return matches[0] if len(matches) == 1 else None
